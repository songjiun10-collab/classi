#!/usr/bin/env python3
"""pdf_to_docx.py — PDF 문제별 캡처 → DOCX

find_problem_boxes로 문제 경계 감지 → 각 문제를 PNG 크롭 → DOCX에 순서대로 삽입.
문제 박스 감지 실패 시 전체 페이지 이미지로 대체.

사용:
  python pdf_to_docx.py exam.pdf
  python pdf_to_docx.py exam.pdf -o output.docx
  python pdf_to_docx.py exam.pdf --zoom 2.5
  python pdf_to_docx.py exam.pdf --pages 1-10
"""

import sys, io, argparse
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core.classifier_engine import find_problem_boxes


def _setup_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(1.5)
    return doc


def _usable_width(doc):
    s = doc.sections[0]
    return s.page_width.cm - s.left_margin.cm - s.right_margin.cm


def _insert_image(doc, img_bytes: bytes, width_cm: float):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after  = Pt(3)
    para.add_run().add_picture(io.BytesIO(img_bytes), width=Cm(width_cm))


def convert(pdf_path: Path, out_path: Path, zoom: float = 2.0, page_range=None):
    doc = _setup_doc()
    usable_w = _usable_width(doc)
    mat = fitz.Matrix(zoom, zoom)

    with fitz.open(pdf_path) as pdf:
        total = pdf.page_count
        pages = range(total) if page_range is None else range(
            max(0, page_range[0] - 1), min(total, page_range[1])
        )
        for i in pages:
            page = pdf[i]
            boxes = find_problem_boxes(page)
            if boxes:
                for num, rect in boxes:
                    try:
                        pix = page.get_pixmap(matrix=mat, clip=rect)
                        _insert_image(doc, pix.tobytes("png"), usable_w)
                    except Exception:
                        pass
            else:
                pix = page.get_pixmap(matrix=mat)
                _insert_image(doc, pix.tobytes("png"), usable_w)
            print(f"  p{i+1}/{total}", end="\r", flush=True)

    print(f"\n저장: {out_path}")
    doc.save(out_path)
    print(f"완료: {out_path.stat().st_size/1024/1024:.1f}MB")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pdf", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=None)
    ap.add_argument("--zoom", type=float, default=2.0)
    ap.add_argument("--pages", type=str, default=None)
    args = ap.parse_args()

    pdf = args.pdf.expanduser().resolve()
    if not pdf.exists():
        sys.exit(f"파일 없음: {pdf}")
    out = args.output or pdf.with_suffix(".docx")
    page_range = None
    if args.pages:
        a, b = args.pages.split("-")
        page_range = (int(a), int(b))
    print(f"\n[{pdf.name}] → {out}")
    convert(pdf, out, args.zoom, page_range)


if __name__ == "__main__":
    main()
